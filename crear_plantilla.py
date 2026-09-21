import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generar_plantilla_excel(ruta_salida="Plantilla_Carga_Forecasting.xlsx"):
    wb = openpyxl.Workbook()
    
    # ---------------------------------------------------------
    # PESTAÑA 1: Carga_Datos
    # ---------------------------------------------------------
    ws_datos = wb.active
    ws_datos.title = "Carga_Datos"
    
    # Definición de columnas
    cols_requeridas = [
        "SKU", "Dias_Operativos_Bloque", "Tiempo_Entrega_Dias", "Nivel_Servicio_Deseado"
    ]
    cols_opcionales = [
        "Descripcion", "Inventario_Actual", "Costo_Unitario", 
        "Tasa_Mantenimiento_Anual", "Costo_Ordenar", "Multiplo_Empaque"
    ]
    cols_historico = [f"M{i:02d}" for i in range(1, 37)]
    
    headers = ["SKU", "Descripcion", "Dias_Operativos_Bloque", "Tiempo_Entrega_Dias", 
               "Nivel_Servicio_Deseado", "Inventario_Actual", "Costo_Unitario", 
               "Tasa_Mantenimiento_Anual", "Costo_Ordenar", "Multiplo_Empaque"] + cols_historico

    ws_datos.append(headers)

    # Datos de ejemplo (SKUs analizados previamente)
    ejemplo_1 = [1394000, "Faro Halógeno H4", 30, 5, 1.65, 12, 6.10, 0.10, 1.00, 50] + [
        0, 0, 0, 0, 0, 241, 322, 296, 246, 339, 390, 225, 794, 297, 282, 422, 
        212, 274, 120, 249, 271, 513, 338, 343, 154, 708, 235, 310, 154, 199, 
        205, 225, 211, 333, 342, 298
    ]
    ejemplo_2 = [1974400, "Faro Incandescente", 30, 5, 1.65, 5, 6.16, 0.10, 1.00, 50] + [
        0, 0, 0, 0, 0, 109, 230, 94, 141, 183, 88, 123, 108, 201, 114, 168, 
        58, 167, 38, 62, 25, 58, 104, 132, 63, 82, 156, 71, 116, 109, 
        48, 123, 140, 196, 87, 67
    ]
    ejemplo_3 = [6203023, "Foco Auxiliar LED", 30, 5, 1.65, 0, 2.50, None, None, None] + [
        0, 0, 0, 0, 0, 1, 1, 4, 7, 12, 10, 7, 2, 9, 10, 0, 0, 5, 2, 29, 
        5, 1, 4, 2, 2, 1, 1, 5, 1, 1, 0, 12, 56, 25, 10, 3
    ]

    ws_datos.append(ejemplo_1)
    ws_datos.append(ejemplo_2)
    ws_datos.append(ejemplo_3)

    # Estilos visuales para encabezados
    fill_req = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")  # Azul Oscuro (Obligatorios)
    fill_opc = PatternFill(start_color="595959", end_color="595959", fill_type="solid")  # Gris (Opcionales)
    fill_hist = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid") # Azul Medio (Histórico)
    
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws_datos.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.alignment = align_center
        
        if col_name in cols_requeridas:
            cell.fill = fill_req
        elif col_name in cols_opcionales:
            cell.fill = fill_opc
        else:
            cell.fill = fill_hist

    # Formatos de celdas de datos
    for row in ws_datos.iter_rows(min_row=2, max_row=4, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Ajustar ancho automático de columnas
    for col in ws_datos.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_datos.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ---------------------------------------------------------
    # PESTAÑA 2: Guia_Campos (Diccionario de Uso)
    # ---------------------------------------------------------
    ws_guia = wb.create_sheet(title="Guia_Campos")
    ws_guia.append(["Campo", "Tipo", "Descripción y Reglas de Negocio"])
    
    diccionario = [
        ("SKU", "Obligatorio", "Código identificador único del artículo (no repetir)."),
        ("Descripcion", "Opcional", "Nombre o detalle comercial del producto."),
        ("Dias_Operativos_Bloque", "Obligatorio", "Días hábiles del periodo base (30 = mes, 15 = quincena, 7 = semana)."),
        ("Tiempo_Entrega_Dias", "Obligatorio", "Lead Time del proveedor en días corridos/operativos."),
        ("Nivel_Servicio_Deseado", "Obligatorio", "Coeficiente Z deseado (1.65 = 95%, 2.05 = 98%, 1.28 = 90%)."),
        ("Inventario_Actual", "Opcional", "Stock disponible. Si se llena, calcula la sugerencia de reabasto neto."),
        ("Costo_Unitario", "Opcional", "Costo de compra unitario. Requerido si se calculará Lote Económico (Q*)."),
        ("Tasa_Mantenimiento_Anual", "Opcional", "Porcentaje decimal (ej. 0.10). Si no se entrega junto con Costo_Ordenar, se omite Q*."),
        ("Costo_Ordenar", "Opcional", "Costo por emisión de orden. Si no se entrega junto con Tasa, se omite Q*."),
        ("Multiplo_Empaque", "Opcional", "Lote o tamaño de caja. Si se omite, los stocks no se redondean a múltiplos."),
        ("M01 a M36", "Obligatorio", "Historial consecutivo de ventas de los últimos 36 periodos (del más antiguo al más reciente).")
    ]
    
    for row in diccionario:
        ws_guia.append(row)
        
    for cell in ws_guia[1]:
        cell.font = font_header
        cell.fill = fill_req
        cell.alignment = align_center

    for col in ws_guia.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_guia.column_dimensions[col_letter].width = max_len + 4

    # Guardar libro
    wb.save(ruta_salida)
    print(f"✅ Plantilla oficial creada exitosamente en: {ruta_salida}")

if __name__ == "__main__":
    generar_plantilla_excel()